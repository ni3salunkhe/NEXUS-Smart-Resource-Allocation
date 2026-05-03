import collections
import collections.abc
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define some generic layouts
    title_slide_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "NEXUS: Smart Resource Allocation"
    subtitle.text = "Unified Operational Nervous System for High-Stakes Field Response\nSolution Challenge 2026"

    # Slide 2: Brief about your solution
    slide = prs.slides.add_slide(content_layout)
    title = slide.shapes.title
    title.text = "Brief about your solution"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "NEXUS is an intelligent, multi-tenant crisis response and resource orchestration platform."
    p = tf.add_paragraph()
    p.text = "It streamlines humanitarian aid by connecting NGOs, field workers, and coordinators in real-time."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Replaces fragmented communication channels with a centralized registry."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Automated data ingestion from unstructured sources (mobile, text, image, CSV)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Features an AI-driven Intelligence Engine to prioritize severe needs and optimally dispatch resources."
    p.level = 1

    # Slide 3: Opportunities
    slide = prs.slides.add_slide(content_layout)
    title = slide.shapes.title
    title.text = "Opportunities"
    content = slide.placeholders[1]
    tf = content.text_frame
    
    tf.text = "a. How different is it from existing ideas?"
    p = tf.add_paragraph()
    p.text = "Moves beyond siloed systems by offering cross-tenant NGO collaboration and multi-modal, unstructured data ingestion (SMS, Images) converted directly into structured task pipelines."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "b. How will it solve the problem?"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "It bridges the critical gap between field discovery and resource dispatch. The Intelligence Engine scores urgency so critical cases are handled first, eliminating operational bottlenecks."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "c. USP of the proposed solution"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "AI-driven intelligent triage, secure DPDP-compliant multi-tenant sharing, and real-time offline-first mobile payload ingestion."
    p.level = 1

    # Slide 4: List of features offered by the solution
    slide = prs.slides.add_slide(content_layout)
    title = slide.shapes.title
    title.text = "List of features offered by the solution"
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Multi-Tenant Household Registry with strict DPDP-compliant consent management."
    p = tf.add_paragraph()
    p.text = "Multi-modal Data Ingestion (Mobile App, WhatsApp/SMS, Images, CSV batch processing)."
    p = tf.add_paragraph()
    p.text = "Real-time Intelligence Engine for Needs Triage and priority algorithmic scoring."
    p = tf.add_paragraph()
    p.text = "Task Coordination & Automated Volunteer Dispatch based on proximity and skills."
    p = tf.add_paragraph()
    p.text = "Live Dashboard Analytics and Heatmaps to identify 'Resource Deserts'."
    p = tf.add_paragraph()
    p.text = "Resilient Architecture with degraded mode operations for unreliable networks."

    # Slide 5: Process flow diagram or Use-case diagram
    slide = prs.slides.add_slide(blank_layout)
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    txBox.text_frame.text = "Process Flow Diagram"
    txBox.text_frame.paragraphs[0].font.size = Pt(32)
    txBox.text_frame.paragraphs[0].font.bold = True
    
    flow_steps = [
        ("1. Data Ingestion", "Field Worker / Beneficiary submits need via App/SMS"),
        ("2. Structuring", "Ingestion Service processes and structures the data"),
        ("3. Triage & Scoring", "Intelligence Engine scores urgency and creates tasks"),
        ("4. Dispatch", "Coordination Service matches task with available volunteer"),
        ("5. Resolution", "Volunteer accepts task, provides relief, and closes loop")
    ]
    
    for i, (title_text, desc_text) in enumerate(flow_steps):
        left = Inches(1)
        top = Inches(1.5 + i * 1.0)
        width = Inches(8)
        height = Inches(0.8)
        
        shape = slide.shapes.add_shape(
            pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, 
            left, top, width, height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(49, 46, 129) # Brand color
        
        tf = shape.text_frame
        tf.text = f"{title_text} : {desc_text}"
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # Slide 6: Architecture diagram
    slide = prs.slides.add_slide(blank_layout)
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1))
    txBox.text_frame.text = "Architecture Diagram of the Proposed Solution"
    txBox.text_frame.paragraphs[0].font.size = Pt(32)
    txBox.text_frame.paragraphs[0].font.bold = True
    
    # Simple Block Diagram for Architecture
    # Frontend
    shape_fe = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.5), Inches(8), Inches(1))
    shape_fe.text_frame.text = "FRONTEND: React + Vite + Zustand (Role-Based Dashboards)"
    
    # API Gateway
    shape_gw = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.8), Inches(8), Inches(0.8))
    shape_gw.text_frame.text = "API GATEWAY & WebSocket Bridge (Real-time events)"
    
    # Microservices Row
    services = ["Auth", "Registry\n(Postgres)", "Ingestion\n(Mongo)", "Intelligence\n(Elastic)", "Coordination\n(Kafka)"]
    for i, svc in enumerate(services):
        s = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(1 + i*1.6), Inches(4), Inches(1.4), Inches(1.2))
        s.text_frame.text = svc
        s.text_frame.paragraphs[0].font.size = Pt(14)
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor(79, 70, 229)
        
    # Infrastructure Row
    shape_infra = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(1), Inches(5.6), Inches(8), Inches(0.8))
    shape_infra.text_frame.text = "INFRASTRUCTURE: Docker, Apache Kafka (Event Bus), Redis (Cache)"
    shape_infra.fill.solid()
    shape_infra.fill.fore_color.rgb = RGBColor(100, 116, 139)

    # Slide 7: Technologies
    slide = prs.slides.add_slide(content_layout)
    title = slide.shapes.title
    title.text = "Technologies to be used in the solution"
    content = slide.placeholders[1]
    tf = content.text_frame
    
    tf.text = "Frontend Ecosystem:"
    p = tf.add_paragraph()
    p.text = "React 19, Vite, TypeScript, TailwindCSS, Zustand (State Management), React Router."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Backend & Microservices:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Python, FastAPI (Asynchronous APIs)."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Databases & Storage:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "PostgreSQL (Relational data), MongoDB (Unstructured data), Elasticsearch (Heatmaps & Text search)."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Messaging & Infrastructure:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Apache Kafka (Event Streaming), Redis (Caching), Docker & Docker Compose."
    p.level = 1

    # Slide 8: Estimated implementation cost
    slide = prs.slides.add_slide(content_layout)
    title = slide.shapes.title
    title.text = "Estimated implementation cost"
    content = slide.placeholders[1]
    tf = content.text_frame
    
    tf.text = "Cloud Hosting (AWS / GCP):"
    p = tf.add_paragraph()
    p.text = "Managed Kubernetes Cluster (EKS/GKE): ~$150/month"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Managed Databases (RDS, MongoDB Atlas, Elastic Cloud): ~$300/month"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Third-Party Integrations:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "SMS & Communication APIs (e.g., Twilio): ~$50/month (Usage based)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Mapping APIs (Mapbox/Google Maps): ~$50/month"
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Software Licensing:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Zero cost - built entirely on open-source technologies (React, Postgres, Kafka)."
    p.level = 1

    prs.save('e:/NEXUS_copy/NEXUS_Solution_Challenge.pptx')

if __name__ == '__main__':
    create_presentation()
